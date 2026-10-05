# X-KR-1 가계부 → 은퇴 나이 엑셀 만들기 (firemap-venture-builder 2026-10-01)
# 손으로 엑셀을 고치지 않는다. 이 스크립트로 다시 만든다.
#   py -3.12 make_xlsx.py            → 가계부_은퇴나이_2026.xlsx (예시 줄 포함)
#   py -3.12 make_xlsx.py --case N   → checks 용 예시 N(1~3) 값으로 만든 파일(out/case_N.xlsx)
#
# 식: src/utils/retirementSimulator.js 의 simulateRetirement() + findEarliestRetirementAge() 를
#     국내 투자(investType 0)·연금 0·부업/임대 0·급여상승 0·90세까지 조건으로 그대로 옮겼다.
#     - 행 A세 = A세가 된 시점 자산. 파이어 나이 R이면 age > R 부터 인출, age <= R(그리고 첫해 아님)이면 저축.
#     - 인출 = 월 생활비 × 12 × (1+물가)^경과년수, 첫해(경과 0)는 수익·저축 없음.
#     - 현금흐름 뒤 자산이 0보다 클 때만 수익률을 곱한다. 은퇴 뒤 자산이 0 이하가 되면 실패.
#     - 은퇴 나이 = 현재 나이~70세 중 90세까지 자산이 안 마르는 가장 이른 R (findEarliestRetirementAge).
#     - 필요 자산 = 지금 그만둬도 90세까지 안 마르는 최소 자산(findRequiredAssetNow, 이분 탐색 대신 닫힌 식).
#   큰 숫자는 웹과 같은 해 단위(R세)다(본부장 판정 2026-10-01, 기준 하나). 'N세 M개월' 어림은 '지난달보다 ±k개월' 줄에만 쓴다.
import sys, os, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import ScatterChart, BarChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'out')

INK, INK2, ORANGE, LINE, BG = '18191D', '4E5968', 'FF5A00', 'E5E8EB', 'F6F7F9'
FONT = '맑은 고딕'
# 통계청 가계동향조사 소비지출 12대 비목(2025년 4/4분기 결과 보도자료 표기, korea.kr newsId=156746265 확인 2026-10-01)
CATS = ['식료품·비주류음료', '주류·담배', '의류·신발', '주거·수도·광열', '가정용품·가사서비스', '보건',
        '교통', '통신', '오락·문화', '교육', '음식·숙박', '기타상품·서비스']
INCOME = '수입'
ROWS = 500           # 가계부 입력 줄 수
FIRST = 3            # 가계부 첫 입력 행(1 머리, 2 머리 설명)
LAST = FIRST + ROWS - 1
KMAX, YMAX = 50, 70  # 계산 격자: 파이어 나이 후보 cur+0..cur+50, 경과 0..70년
UNTIL, MAXR = 90, 70
DELTA = 100000       # '+N일' 기울기용 월 지출 증가분(원)

EX_LINES = [  # (항목, 금액, 분류) — 10월 예시
    ('급여', 3500000, INCOME), ('월세·관리비', 700000, '주거·수도·광열'), ('장보기', 400000, '식료품·비주류음료'),
    ('외식·배달', 300000, '음식·숙박'), ('커피', 100000, '음식·숙박'), ('교통카드', 120000, '교통'),
    ('휴대폰', 80000, '통신'), ('OTT·취미', 150000, '오락·문화'), ('기타', 150000, '기타상품·서비스')]
CASES = {  # checks.md 예시: (나이, 자산, 수익률%, 물가%, 월 수입, 이번 달 지출, 지난달 지출)
    1: (30, 30000000, 5, 3, 3500000, 2000000, 2150000),
    2: (35, 100000000, 5, 3, 5000000, 3000000, 3000000),
    3: (40, 200000000, 5, 3, 4500000, 2500000, 2800000),
}


def f(size=11, bold=False, color=INK):
    return Font(name=FONT, size=size, bold=bold, color=color)


def ledger_rows(case):
    """(날짜, 항목, 금액, 분류) 목록. 예시 파일은 9·10월 두 달."""
    if case is None:
        sep = [(n, a + (100000 if n == '외식·배달' else 50000 if n == '기타' else 0), c) for n, a, c in EX_LINES]
        out = [(datetime.date(2026, 9, 1 + i), n, a, c) for i, (n, a, c) in enumerate(sep)]
        out += [(datetime.date(2026, 10, 1 + i), n, a, c) for i, (n, a, c) in enumerate(EX_LINES)]
        return out
    age, asset, r, i, inc, m_this, m_prev = CASES[case]
    return [(datetime.date(2026, 9, 1), '급여', inc, INCOME), (datetime.date(2026, 9, 2), '생활비', m_prev, '기타상품·서비스'),
            (datetime.date(2026, 10, 1), '급여', inc, INCOME), (datetime.date(2026, 10, 2), '생활비', m_this, '기타상품·서비스')]


def build(case=None):
    wb = Workbook()
    s1 = wb.active
    s1.title = '월 가계부'
    s2 = wb.create_sheet('은퇴 나이')
    s3 = wb.create_sheet('설명')
    c = wb.create_sheet('계산')
    c.sheet_state = 'hidden'
    thin = Side(style='thin', color=LINE)

    # ── 시트 1: 날짜 · 항목 · 금액 · 분류 · 은퇴 +N일 (카드 내역 앞 세 열 붙여넣기용)
    for col, (h, w) in enumerate([('날짜', 12), ('항목', 22), ('금액', 14), ('분류', 20), ('은퇴 +N일', 12)], 1):
        s1.cell(1, col, h).font = f(11, True)
        s1.column_dimensions[L(col)].width = w
        s1.cell(1, col).border = Border(bottom=thin)
    s1.cell(2, 5, '매달 반복된다고 가정할 때').font = f(9, color=INK2)
    s1.freeze_panes = 'A3'
    dv = DataValidation(type='list', formula1='=계산!$H$2:$H$14', allow_blank=True)
    s1.add_data_validation(dv)
    dv.add(f'D{FIRST}:D{LAST}')
    for r in range(FIRST, LAST + 1):
        s1.cell(r, 1).number_format = 'yyyy-mm-dd'
        s1.cell(r, 3).number_format = '#,##0'
        e = s1.cell(r, 5, f'=IF(OR(C{r}="",D{r}="",D{r}="{INCOME}",계산!$B$20=""),"",ROUND(C{r}*계산!$B$20,0))')
        e.number_format = '"+"#,##0"일";"−"#,##0"일";"0일"'
        e.font = f(11, True, ORANGE)
    for k, (d, n, a, cat) in enumerate(ledger_rows(case)):
        r = FIRST + k
        s1.cell(r, 1, d); s1.cell(r, 2, n); s1.cell(r, 3, a); s1.cell(r, 4, cat)
    # 월 요약 G:K (최근 12개월)
    for col, (h, w) in enumerate([('월', 10), ('수입', 13), ('지출', 13), ('저축', 13), ('저축률', 9)], 7):
        s1.cell(1, col, h).font = f(11, True)
        s1.cell(1, col).border = Border(bottom=thin)
        s1.column_dimensions[L(col)].width = w
    s1.column_dimensions['F'].width = 3
    A, C, D = f'$A${FIRST}:$A${LAST}', f'$C${FIRST}:$C${LAST}', f'$D${FIRST}:$D${LAST}'
    for k in range(12):
        r = 3 + k
        s1.cell(r, 7, f'=IF(계산!$B$5="","",EDATE(계산!$B$5,{k - 11}))').number_format = 'yyyy-mm'
        rng = f'{A},">="&G{r},{A},"<"&EDATE(G{r},1)'
        s1.cell(r, 8, f'=IF(G{r}="","",SUMIFS({C},{D},"{INCOME}",{rng}))').number_format = '#,##0;-#,##0;""'
        s1.cell(r, 9, f'=IF(G{r}="","",SUMIFS({C},{rng})-H{r})').number_format = '#,##0;-#,##0;""'  # 빈 달은 0 대신 빈칸
        s1.cell(r, 10, f'=IF(OR(G{r}="",AND(H{r}=0,I{r}=0)),"",H{r}-I{r})').number_format = '#,##0'
        s1.cell(r, 11, f'=IF(OR(G{r}="",H{r}=0),"",J{r}/H{r})').number_format = '0%'
        for col in range(7, 12):
            s1.cell(r, col).font = f(10)

    s1.print_area = 'A1:K60'  # 수식이 든 빈 줄 500개까지 인쇄되지 않게
    s1.page_setup.fitToWidth = 1
    s1.page_setup.fitToHeight = 0
    s1.sheet_properties.pageSetUpPr.fitToPage = True

    # ── 시트 2: 은퇴 나이 (spec ⓑ)
    s2.sheet_view.showGridLines = False
    s2.column_dimensions['A'].width = 2
    for col in 'BCDEF':
        s2.column_dimensions[col].width = 14
    s2.column_dimensions['G'].width = 6
    # spec ⓑ의 B2 시트 제목은 B3 라벨과 같은 말이라 두 번 보인다(PDF 실측) → 시트 탭 이름으로 대신, B3 라벨만 둔다
    s2['B3'] = '은퇴 나이'; s2['B3'].font = f(11, color=INK2)
    s2.merge_cells('B4:D6')
    s2['B4'] = '=계산!B30'
    s2['B4'].font = f(36, True, ORANGE)
    s2['B4'].alignment = Alignment(vertical='center')
    s2['B7'] = '=계산!B31'; s2['B7'].font = f(11, color=INK2)
    for col, lab, ref, fmt in [('B', '연 지출', '=계산!B6*12', '#,##0"원"'), ('D', '필요 자산', '=계산!B32', '#,##0"원"'),
                               ('F', '현재 자산', '=D13', '#,##0"원"')]:
        s2[f'{col}9'] = lab; s2[f'{col}9'].font = f(10, color=INK2)
        s2[f'{col}10'] = ref; s2[f'{col}10'].font = f(14, True); s2[f'{col}10'].number_format = fmt
        nxt = chr(ord(col) + 1)
        s2.merge_cells(f'{col}10:{nxt}10')  # 9자리 원 금액이 14pt에서 한 칸에 안 들어간다(#### 실측)
        for cc in (col, nxt):
            s2[f'{cc}10'].border = Border(bottom=thin)
    inputs = [('현재 나이', 30, '세', '0'), ('현재 자산', 30000000, '원', '#,##0'), ('연 수익률', 5, '%', '0.0'), ('물가', 3, '%', '0.0')]
    if case:
        age, asset, rr, ii = CASES[case][:4]
        inputs = [(a, v, u, fm) for (a, _, u, fm), v in zip(inputs, (age, asset, rr, ii))]
    for k, (lab, val, unit, fmt) in enumerate(inputs):
        r = 12 + k
        s2[f'B{r}'] = lab; s2[f'B{r}'].font = f(11)
        s2.merge_cells(f'D{r}:F{r}')
        s2[f'D{r}'] = val; s2[f'D{r}'].number_format = fmt; s2[f'D{r}'].font = f(11, True)
        for col in 'DEF':
            s2[f'{col}{r}'].fill = PatternFill('solid', fgColor=BG)
            s2[f'{col}{r}'].border = Border(top=thin, bottom=thin, left=thin if col == 'D' else None, right=thin if col == 'F' else None)
        s2[f'G{r}'] = unit; s2[f'G{r}'].font = f(11, color=INK2)
    s2['B18'] = '파이어맵 계산기와 같은 식 · firemap.kr'
    s2['B18'].hyperlink = 'https://firemap.kr/?utm_source=xlsx&utm_medium=sheet&utm_campaign=xkr1'
    s2['B18'].font = Font(name=FONT, size=10, color=INK2, underline='single')
    s2.print_area = 'A1:H50'
    s2.page_setup.orientation = 'landscape'
    s2.page_setup.fitToWidth = 1
    s2.sheet_properties.pageSetUpPr.fitToPage = True

    # ── 시트 3: 설명 (편집 통과 2026-10-01 firemap-editor-web)
    s3.sheet_view.showGridLines = False
    s3.column_dimensions['A'].width = 2
    s3.column_dimensions['B'].width = 100
    lines = [
        ('쓰는 법', True),
        ('1. 「월 가계부」에 날짜·항목·금액을 넣고 분류를 목록에서 고릅니다. 월급 같은 들어온 돈은 분류를 「수입」으로 둡니다.', False),
        ('2. 카드사·가계부 앱에서 내려받은 내역의 날짜·내용·금액 세 열을 A~C열에 그대로 붙여 넣을 수 있습니다. 열 순서가 다르면 순서만 맞춰 주세요.', False),
        ('3. 지출 줄 오른쪽 「은퇴 +N일」은 그 지출이 매달 반복된다고 가정할 때 은퇴 나이가 며칠 늦어지는지입니다.', False),
        ('4. 「은퇴 나이」 시트에 현재 나이·현재 자산·연 수익률·물가 네 칸을 넣으면 가장 최근 달의 수입·지출로 은퇴 나이를 계산합니다.', False),
        ('5. 예시 줄(2026년 9~10월)은 지우고 쓰세요. 다음 달에 줄을 더하면 「지난달보다」가 바뀝니다.', False),
        ('', False),
        ('계산 가정', True),
        ('· 최근 달 (수입 − 지출)을 은퇴할 때까지 매달 모으고, 은퇴 뒤에는 그달 지출을 매년 물가만큼 늘려 꺼내 씁니다.', False),
        ('· 90세까지 자산이 남는 가장 이른 나이를 찾습니다(70세까지). 국민연금·퇴직금·부동산·세금은 넣지 않았습니다.', False),
        ('· 「지난달보다 N개월 빨라져요·늦어져요」는 앞뒤 두 해 결과 사이를 나눠 어림한 값입니다.', False),
        ('· 「나이별 자산」 그래프는 은퇴 나이에 그만둘 때 90세까지 자산이 어떻게 바뀌는지, 「월별 저축률」은 최근 12개월 (수입 − 지출) ÷ 수입입니다.', False),
        ('· 「은퇴 +N일」은 월 지출이 10만원 늘 때 바뀌는 날 수에 비례해 어림한 값입니다.', False),
        ('', False),
        ('알림', True),
        ('· 이 파일은 계산 도구이며 투자 조언이 아닙니다. 특정 종목·금융상품을 추천하지 않고, 수익을 보장하지 않습니다.', False),
        ('· 지출 분류: 통계청 가계동향조사 소비지출 12대 비목.', False),
        ('· 기준일 2026-10-01', False),
        ('firemap.kr', False),
    ]
    for k, (t, bold) in enumerate(lines):
        cell = s3.cell(2 + k, 2, t)
        cell.font = f(12 if bold else 11, bold)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    link = s3.cell(2 + len(lines) - 1, 2)
    link.hyperlink = 'https://firemap.kr/?utm_source=xlsx&utm_medium=guide&utm_campaign=xkr1'
    link.font = Font(name=FONT, size=11, color=INK2, underline='single')

    # ── 계산(숨김)
    c['A1'] = '나이'; c['B1'] = "='은퇴 나이'!D12"
    c['A2'] = '자산'; c['B2'] = "='은퇴 나이'!D13"
    c['A3'] = '수익률'; c['B3'] = "='은퇴 나이'!D14/100"
    c['A4'] = '물가'; c['B4'] = "='은퇴 나이'!D15/100"
    c['A5'] = '최근 달'; c['B5'] = f"=IF(COUNT('월 가계부'!$A${FIRST}:$A${LAST})=0,\"\",DATE(YEAR(MAX('월 가계부'!$A${FIRST}:$A${LAST})),MONTH(MAX('월 가계부'!$A${FIRST}:$A${LAST})),1))"
    c['A6'] = '이번 달 지출'; c['B6'] = "=IF(B5=\"\",0,INDEX('월 가계부'!$I$3:$I$14,12))"
    c['A7'] = '이번 달 저축'; c['B7'] = "=IF(B5=\"\",0,INDEX('월 가계부'!$J$3:$J$14,12))"
    c['A8'] = '지난달 지출'; c['B8'] = "=IF(B5=\"\",0,INDEX('월 가계부'!$I$3:$I$14,11))"
    c['A9'] = '지난달 저축'; c['B9'] = "=IF(B5=\"\",0,INDEX('월 가계부'!$J$3:$J$14,11))"
    c['A10'] = '기울기 증가분'; c['B10'] = DELTA
    c['H1'] = '분류 목록'
    for k, cat in enumerate([INCOME] + CATS):
        c.cell(2 + k, 8, cat)
    # 시나리오: 0 이번 달 · 1 이번 달 + 월 지출 10만 · 2 지난달
    scen = [('B6', 'B7'), ('(B6+B10)', '(B7-B10)'), ('B8', 'B9')]
    tops = []
    for s, (M, S) in enumerate(scen):
        top = 40 + s * (KMAX + 6)
        tops.append(top)
        c.cell(top - 1, 1, f'시나리오 {s}')
        c.cell(top - 1, 2, 'R')
        for y in range(YMAX + 1):
            c.cell(top - 1, 3 + y, y)
        fin, sur, fr = 3 + YMAX + 1, 3 + YMAX + 2, 3 + YMAX + 3
        c.cell(top - 1, fin, '90세 자산'); c.cell(top - 1, sur, '가능 R')
        Mx, Sx = M.replace('B', '$B$'), S.replace('B', '$B$')
        for k in range(KMAX + 1):
            r = top + k
            c.cell(r, 1, k)
            c.cell(r, 2, f'=$B$1+{k}')
            c.cell(r, 3, '=$B$2')
            for y in range(1, YMAX + 1):
                p = f'{L(2 + y)}{r}'
                flow = f'IF($B$1+{y}<=$B{r},12*{Sx},-12*{Mx}*(1+$B$4)^{y})'
                c.cell(r, 3 + y, f'=IF($B$1+{y}>{UNTIL},{p},({p}+{flow})*IF({p}+{flow}>0,1+$B$3,1))')
            c.cell(r, fin, f'={L(3 + YMAX)}{r}')
            c.cell(r, sur, f'=IF(AND($B{r}<={MAXR},{L(fin)}{r}>0),$B{r},999)')
        rngR = f'{L(sur)}{top}:{L(sur)}{top + KMAX}'
        rngF = f'{L(fin)}{top}:{L(fin)}{top + KMAX}'
        # 결과 칸: B(top-3)=가장 이른 R(정수, 웹과 같음), C=어림 나이(소수), D=개월 합
        hr = top - 3
        c.cell(hr, 1, f'결과 {s}')
        c.cell(hr, 2, f'=MIN({rngR})')
        c.cell(hr, 3, f'=IF(B{hr}=999,"",IF(B{hr}=$B$1,$B$1,B{hr}-1+(-INDEX({rngF},B{hr}-$B$1))/(INDEX({rngF},B{hr}-$B$1+1)-INDEX({rngF},B{hr}-$B$1))))')
        c.cell(hr, 4, f'=IF(C{hr}="","",ROUND(C{hr}*12,0))')
    h0, h1, h2 = (t - 3 for t in tops)
    c['A20'] = '+N일 기울기(일/원)'
    c['B20'] = f'=IF(OR(C{h0}="",C{h1}=""),"",(C{h1}-C{h0})*365/$B$10)'
    c['A21'] = '웹과 같은 해 단위 나이'; c['B21'] = f'=IF(B{h0}=999,"",B{h0})'
    c['A30'] = '큰 숫자'
    c['B30'] = (f'=IF(B5="","가계부를 넣어 주세요",IF(D{h0}="","70세 넘음",IF(B{h0}=$B$1,"지금",'
                f'B{h0}&"세")))')
    c['A31'] = '지난달 대비'
    c['B31'] = (f'=IF(OR(B8=0,D{h0}="",D{h2}=""),"",IF(D{h0}=D{h2},"지난달과 같음",'
                f'"이번 달 지출대로면 지난달보다 "&ABS(D{h0}-D{h2})&"개월 "&IF(D{h0}<D{h2},"빨라져요","늦어져요")))')  # 순돌이 10/5: 무엇과 비교인지
    # 필요 자산: 지금 그만둬도 90세까지 안 마름 ⇔ 자산 > Σ_{y=1..n} 12M(1+i)^y/(1+r)^(y-1), n = 90 - 나이
    c['A32'] = '필요 자산'
    c['CB1'] = 'y'; c['CC1'] = '할인 인출'  # 격자(A~BX)와 겹치지 않는 자리
    for y in range(1, YMAX + 1):
        c.cell(1 + y, 80, y)
        c.cell(1 + y, 81, f'=IF($B$1+CB{1 + y}>{UNTIL},0,12*$B$6*(1+$B$4)^CB{1 + y}/(1+$B$3)^(CB{1 + y}-1))')
    c['B32'] = f'=CEILING(SUM(CC2:CC{1 + YMAX}),1)'

    # ── 그래프 2개(순돌이 10/5 X-KR-1 ①: 경쟁 4번도 차트) — 시트 2 아래, 숫자는 계산 시트 값 그대로
    # 나이별 자산: 시나리오 0 격자에서 은퇴 나이 R 행을 그대로 읽는다(90세 넘으면 NA → 선이 끊김)
    c['CE1'] = '나이'; c['CF1'] = '자산'
    g0 = tops[0]
    for y in range(YMAX + 1):
        r = 2 + y
        c.cell(r, 83, f'=IF($B$1+{y}>{UNTIL},NA(),$B$1+{y})')
        c.cell(r, 84, f'=IF(OR($B$21="",$B$1+{y}>{UNTIL}),NA(),INDEX($C${g0}:${L(3 + YMAX)}${g0 + KMAX},$B$21-$B$1+1,{y + 1}))')
    ch = ScatterChart()
    ch.title = '나이별 자산(백만원) — 은퇴 나이에 그만둘 때'
    ch.style = 2
    ch.y_axis.number_format = '#,##0,,'
    ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = 20, 90, 10
    ch.x_axis.majorGridlines = None
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    ch.legend = None
    ch.visible_cells_only = False  # 숨긴 계산 시트 값도 그린다
    se = Series(Reference(c, min_col=84, min_row=2, max_row=2 + YMAX), Reference(c, min_col=83, min_row=2, max_row=2 + YMAX))
    se.marker.symbol = 'none'
    se.smooth = False
    se.graphicalProperties.line.solidFill = ORANGE
    se.graphicalProperties.line.width = 28000
    ch.series.append(se)
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.height, ch.width = 7.5, 16
    s2.add_chart(ch, 'B21')
    bc = BarChart()
    bc.title = '월별 저축률'
    bc.style = 2
    bc.y_axis.number_format = '0%'
    bc.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    bc.legend = None
    bc.add_data(Reference(s1, min_col=11, min_row=3, max_row=14), titles_from_data=False)
    bc.set_categories(Reference(s1, min_col=7, min_row=3, max_row=14))
    bc.series[0].graphicalProperties.solidFill = INK2
    bc.x_axis.number_format = 'yy-mm'
    bc.x_axis.delete = False; bc.y_axis.delete = False
    bc.height, bc.width = 6.5, 16
    s2.add_chart(bc, 'B37')
    return wb


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    if '--case' in sys.argv:
        n = int(sys.argv[sys.argv.index('--case') + 1])
        p = os.path.join(OUT, f'case_{n}.xlsx')
    else:
        n, p = None, os.path.join(OUT, '가계부_은퇴나이_2026.xlsx')
    build(n).save(p)
    print(p)
