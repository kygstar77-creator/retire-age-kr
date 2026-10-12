# 카페 회원 자료 1호 「금융소득 2천만원·건보료 1천만원 판정 엑셀」 만들기 (자료 담당, 2026-10-12)
# 손으로 엑셀을 고치지 않는다. 이 스크립트로 다시 만든다.
#   py -3.12 make_fin_xlsx.py        → out/금융소득_건보료_판정_2026.xlsx
# 근거: work/research/growth/calc-gap-1012.md 1순위. 기준값 원문은 fin-excel-check.md 1장.
# 식: 세금은 소득세법 제14조·제17조·제55조·제56조·제62조·제15조, 시행령 제116조의2, 국세상담센터 '금융소득 종합과세 세액계산' 산식.
#     건보는 국민건강보험법 제69조⑤·제71조②, 시행령 제41조·제44조·별표 4, 시행규칙 제44조·별표 1의2, 고시 제2025-222호,
#     노인장기요양보험법 제9조①·시행령 제4조. 사이트 식(src/firemap-v2/healthInsurance.js)과 다른 곳은 fin-excel-check.md 2장.
# 확인 안 된 값은 '기준값' 시트에 '확인 필요'로만 적고 계산에 쓰지 않는다.
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'out')
FILE = os.path.join(OUT, '금융소득_건보료_판정_2026.xlsx')

INK, INK2, ORANGE, LINE, BG = '18191D', '4E5968', 'FF5A00', 'E5E8EB', 'F6F7F9'   # 가계부 엑셀(x-kr-1/make_xlsx.py)과 같은 색
FONT = '맑은 고딕'
DEP, LOC = '피부양자', '지역가입자'
EXAMPLE = dict(interest=8000000, div_kr=10000000, div_etc=6000000, status=DEP, pension=12000000, other=0, prop=300000000)

# ── 기준값(값, 단위, 근거, 원문 문장, 시행·적용, 확인) ─────────────────────────────
# 원문은 law.go.kr DRF API(work/lawtext.py, 2026-10-12 조회: 소득세법 2026.4.21 공포·시행판, 국민건강보험법 시행령 2026.10.1 시행판,
# 시행규칙 2026.9.28 시행판, 노인장기요양보험법 2026.5.26판·시행령 2026.5.12판)와 공단·국세청 누리집에서 옮겼다.
K_ROWS = [
    ('T', '종합과세기준금액', 20000000, '원', '소득세법 제14조제3항제6호',
     '이자소득과 배당소득으로서 그 소득의 합계액이 2천만원(이하 "이자소득등의 종합과세기준금액"이라 한다) 이하이면서 제127조에 따라 원천징수된 소득',
     '현행(2026.4.21 시행판)'),
    ('WI', '원천징수세율 · 이자', 0.14, '', '소득세법 제129조제1항제1호라목', '그 밖의 이자소득에 대해서는 100분의 14', '현행(2026.4.21 시행판)'),
    ('WD', '원천징수세율 · 배당', 0.14, '', '소득세법 제129조제1항제2호나목', '그 밖의 배당소득에 대해서는 100분의 14', '현행(2026.4.21 시행판)'),
    ('GU', '배당가산율(Gross-up)', 0.10, '', '소득세법 제17조제3항, 부칙(법률 제21221호, 2025.12.23) 제14조',
     '해당 과세기간의 총수입금액에 그 배당소득의 100분의 10에 해당하는 금액을 더한 금액 / 2026년 12월 31일 이전에 지급받은 배당소득의 소득금액 계산에 관하여는 … 종전의 규정에 따른다',
     '2026년 12월 31일까지 받는 배당. 2027.1.1부터 100분의 11(같은 부칙 제1조제2호) — 2027년 배당이면 이 칸을 11%로 바꾼다'),
    ('DED', '종합소득공제(본인 기본공제만)', 1500000, '원', '소득세법 제50조제1항제1호',
     '1명당 연 150만원을 곱하여 계산한 금액을 그 거주자의 해당 과세기간의 종합소득금액에서 공제한다',
     '현행. 배우자·부양가족 공제나 다른 소득공제가 있으면 이 칸을 그 합계로 바꾼다'),
    ('HR', '건강보험료율', 0.0719, '', '국민건강보험법 시행령 제44조제1항', '지역가입자의 보험료율은 각각 1만분의 719로 한다',
     '2026년 1월분 보험료부터(시행령 부칙 2025.12.23 제2조). 2027년도 7.19% 유지 심의(정책브리핑 2026.9.8 제15차 건정심 브리핑)'),
    ('PT', '재산보험료 부과점수당 금액', 211.5, '원', '국민건강보험법 시행령 제44조제2항', '지역가입자의 재산보험료부과점수당 금액은 211.5원으로 한다',
     '2026년 1월분 보험료부터(같은 부칙 제2조). 2027년 값은 확인 안 함'),
    ('LTCR', '장기요양보험료율', 0.009448, '', '노인장기요양보험법 시행령 제4조', '장기요양보험료율은 100만분의 9,448로 한다', '2026.5.12판 기준'),
    ('LTC', '장기요양보험료 = 건보료 × 이 비율', '=ROUND({LTCR}/{HR},4)', '', '노인장기요양보험법 제9조제1항',
     '산정한 보험료액에서 … 건강보험료율 대비 장기요양보험료율의 비율(소수점 이하 다섯째자리에서 반올림한다)을 곱하여 산정한 금액 / 보건복지부 보도자료(2025.11.4): "2026년도 건강보험료 대비 장기요양보험료율은 13.14%"', '2026.5.26판'),
    ('FT', '건보 금융소득 문턱', 10000000, '원', '국민건강보험법 시행규칙 제44조제1항 단서',
     '「소득세법」 제14조제3항제6호에 따른 소득이 1천만원 이하인 경우에는 해당 이자소득과 배당소득은 합산하지 않는다', '2026.9.28 시행판'),
    ('HALF', '지역가입자 근로·연금소득 반영률', 0.5, '', '국민건강보험법 시행규칙 제44조제2항제2호·제3항',
     '영 제41조제1항제4호 및 제5호의 소득: 해당 소득의 100분의 50 / 지역가입자의 소득월액 산정방법에 관하여는 제2항 전단을 준용한다', '2026.9.28 시행판'),
    ('FLOORM', '소득월액 하한 기준', 280000, '원', '국민건강보험법 시행규칙 제44조제3항 단서',
     '지역가입자의 소득월액이 28만원 이하인 경우에는 … 월별 보험료액의 하한액을 … 보험료율로 나누어 얻은 값을 해당 지역가입자의 소득월액으로 한다', '2026.9.28 시행판'),
    ('FLOOR', '지역가입자 월별 보험료 하한', 20160, '원', '월별 건강보험료액의 상한과 하한에 관한 고시(보건복지부고시 제2025-222호) 제3조제2호',
     '지역가입자의 월별 보험료액 : 20,160원', '2026년 1월분 보험료부터(고시 부칙)'),
    ('CAP', '지역가입자 월별 보험료 상한', 4591740, '원', '같은 고시 제2조제2호',
     '직장가입자의 소득월액 보험료 및 지역가입자의 월별 보험료액 : 4,591,740원', '2026년 1월분 보험료부터(고시 부칙)'),
    ('PDED', '재산 공제', 100000000, '원', '국민건강보험법 시행령 별표 4 제1호',
     '다음 각 목의 금액을 합산한 금액에서 1억원 및 제42조의2제3항 각 호의 구분에 따른 금액을 뺀 금액을 등급별로 구분하여 산정한다', '2026.10.1 시행판'),
    ('DEPINC', '피부양자 소득 한도', 20000000, '원', '국민건강보험법 시행규칙 별표 1의2 제1호가목',
     '영 제41조제1항 각 호의 소득을 같은 조 제3항의 소득자료에 따라 산정한 합계액이 연간 2천만원 이하일 것', '2026.9.28 시행판'),
    ('DEPP1', '피부양자 재산 기준 1', 540000000, '원', '같은 별표 제2호가목2)', '재산세 과세표준의 합이 5억4천만원 이하일 것', '2026.9.28 시행판'),
    ('DEPP2', '피부양자 재산 기준 2', 900000000, '원', '같은 별표 제2호가목1)',
     '재산세 과세표준의 합이 5억 4천만원을 초과하면서 9억원 이하이고, 영 제41조제1항 각 호에 따른 소득의 합계액이 연간 1천만원 이하일 것', '2026.9.28 시행판'),
    ('DEPINC2', '재산 5.4억~9억일 때 소득 한도', 10000000, '원', '같은 별표 제2호가목1)', '(위 문장) 소득의 합계액이 연간 1천만원 이하일 것', '2026.9.28 시행판'),
]
TAX_TABLE = [  # 소득세법 제55조제1항 표: (초과 기준, 기본 세액, 세율)
    (0, 0, 0.06), (14000000, 840000, 0.15), (50000000, 6240000, 0.24), (88000000, 15360000, 0.35),
    (150000000, 37060000, 0.38), (300000000, 94060000, 0.40), (500000000, 174060000, 0.42), (1000000000, 384060000, 0.45)]
# 국민건강보험법 시행령 별표 4 제3호 재산등급별 점수(재산금액 만원 '초과 ~ 이하') — 2026.10.1 시행판 원문 그대로
PROP_UPPER_MAN = [450, 900, 1350, 1800, 2250, 2700, 3150, 3600, 4050, 4500, 5020, 5590, 6220, 6930, 7710, 8590, 9570, 10700, 11900, 13300,
                  14800, 16400, 18300, 20400, 22700, 25300, 28100, 31300, 34900, 38800, 43200, 48100, 53600, 59700, 66500, 74000, 82400,
                  91800, 103000, 114000, 127000, 142000, 158000, 176000, 196000, 218000, 242000, 270000, 300000, 330000, 363000, 399300,
                  439230, 483153, 531468, 584615, 643077, 707385, 778124]
PROP_POINTS = [22, 44, 66, 97, 122, 146, 171, 195, 219, 244, 268, 294, 320, 344, 365, 386, 412, 439, 465, 490, 516, 535, 559, 586, 611, 637,
               659, 681, 706, 731, 757, 785, 812, 841, 881, 921, 961, 1001, 1041, 1091, 1141, 1191, 1241, 1291, 1341, 1391, 1451, 1511, 1571,
               1641, 1711, 1781, 1851, 1921, 1991, 2061, 2131, 2201, 2271, 2341]
UNCHECKED = [  # (항목, 확인한 것, 이 엑셀에서)
    ('피부양자 소득 2천만원에 1천만원 이하 금융소득을 넣는지',
     '별표 1의2는 "영 제41조제1항 각 호의 소득 … 합계액"이라고만 적었다. 공단 누리집(피부양자 자격 인정기준)은 "모든 소득(사업·금융·연금·근로기타소득 등)을 합하여 연간 2천만원 이하". 1천만원 이하 합산 제외(시행규칙 제44조제1항 단서)는 "소득월액 산정에 포함되는 소득" 조항이라 피부양자 판정에도 쓰는지 원문으로 확인 못 함',
     '두 경우를 다 계산한다. 결과가 같으면 그대로, 다르면 판정 칸에 "확인 필요"'),
    ('건보 소득에 배당가산액을 넣는지', '시행규칙 제44조제1항제2호는 "「소득세법」 제17조제3항에 따라 산정한 배당소득금액". 공단이 가산액을 더하는지 공단 자료로 확인 못 함',
     '계산에 쓰지 않음. 배당은 받은 금액(세전) 그대로 건보 소득에 넣는다'),
    ('공단 지역보험료 모의계산이 1천만원 이하 금융소득을 스스로 빼는지', '계산은 공단 서버에서 하고, 화면은 이자·배당을 "사업소득 등" 칸에 합쳐 넣게 한다. 서버에 계산을 보내 보지 않았다',
     '법 원문(시행규칙 제44조제1항 단서)대로 1천만원 이하는 뺀다'),
    ('국내에서 원천징수되지 않은 이자·배당(해외 계좌로 직접 받은 것 등)', '국세상담센터: "원천징수되지 않은 금융소득(예시 : 국외원천 이자·배당소득)과 출자공동사업자의 배당소득은 2천만원 이하라도 종합과세합니다"',
     '계산하지 않음. 넣으면 2천만원 이하에서 추가 세금이 0으로 잘못 나온다'),
    ('외국납부세액공제', '소득세법 제57조(공제한도금액 = 산출세액 × 국외원천소득 ÷ 종합소득금액)', '계산에 쓰지 않음. 해외주식 배당이 있으면 실제 추가 세금은 이보다 적을 수 있다'),
    ('지방소득세', '지방세법 제93조제2항이 같은 비교 방식(세율은 소득세의 100분의 10 꼴)을 둔다. 배당세액공제에 해당하는 지방세 공제는 확인 못 함', '계산에 쓰지 않음(소득세만)'),
    ('고배당기업 배당 분리과세 특례', '조세특례제한법 제104조의27(2028년 12월 31일이 속하는 사업연도까지 발생하는 특례배당소득)', '계산에 쓰지 않음. 특례배당소득은 넣지 않는다'),
    ('표준세액공제 등 다른 세액공제', '소득세법 제15조제2호: 비교한 큰 금액에서 제56조의2·제57조·제59조의4 등의 세액공제를 한다', '계산에 쓰지 않음(두 세액 모두 공제 전)'),
    ('2027년 이후 점수당 금액·장기요양보험료율·하한·상한', '2027년 보험료율 7.19% 유지 심의만 확인(정책브리핑 2026.9.8)', '2026년 값으로 계산'),
    ('보험료 끝수(10원 미만) 처리', '원문 확인 못 함', '원 단위 반올림. 고지서와 끝자리가 다를 수 있다'),
    ('세대 단위·주택금융부채 공제·전월세', '국민건강보험법 제69조제5항 후단 "보험료액은 세대 단위로 산정한다", 별표 4 제1호나목(전월세)·제42조의2(부채 공제)', '계산하지 않음. 한 사람 소득·재산만 넣는다'),
]


def f(size=11, bold=False, color=INK, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, underline=underline)


def build(path=FILE, ex=EXAMPLE):
    wb = Workbook()
    s = wb.active
    s.title = '판정'
    k = wb.create_sheet('기준값')
    g = wb.create_sheet('설명')
    thin = Side(style='thin', color=LINE)

    # ── 기준값 시트 ───────────────────────────────────────────────
    k.sheet_view.showGridLines = False
    for col, w in zip('ABCDEFGH', (2, 30, 14, 6, 34, 70, 40, 10)):
        k.column_dimensions[col].width = w
    k['B2'] = '기준값 — 해마다 바뀌는 값은 여기만 고친다'; k['B2'].font = f(13, True)
    k['B3'] = '원문은 국가법령정보센터(law.go.kr)·국민건강보험공단·국세청 누리집에서 2026-10-12에 확인했다.'; k['B3'].font = f(10, color=INK2)
    hdr = ['항목', '값', '단위', '근거', '원문 문장', '시행·적용', '확인']
    for i, h in enumerate(hdr):
        c = k.cell(5, 2 + i, h); c.font = f(10, True, INK2); c.border = Border(bottom=thin)
    A = {}
    r = 6
    for key, name, val, unit, src, text, when in K_ROWS:
        A[key] = f"기준값!$C${r}"
        r += 1
    r = 6
    for key, name, val, unit, src, text, when in K_ROWS:
        k.cell(r, 2, name).font = f(10)
        v = val.format(**{x: A[x] for x in A}) if isinstance(val, str) else val
        c = k.cell(r, 3, v); c.font = f(10, True); c.fill = PatternFill('solid', fgColor=BG)
        c.number_format = {'WI': '0%', 'WD': '0%', 'GU': '0%', 'HALF': '0%', 'HR': '0.00%', 'LTCR': '0.0000%', 'LTC': '0.0000', 'PT': '0.0'}.get(key, '#,##0')
        k.cell(r, 4, unit).font = f(10, color=INK2)
        k.cell(r, 5, src).font = f(10)
        k.cell(r, 6, text).font = f(9, color=INK2)
        k.cell(r, 7, when).font = f(9, color=INK2)
        k.cell(r, 8, '확인함').font = f(9, color=INK2)
        for col in range(2, 9):
            k.cell(r, col).alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    # 세율표
    r += 1
    k.cell(r, 2, '종합소득세 기본세율 — 소득세법 제55조제1항(현행, 2026.4.21 시행판)').font = f(11, True); r += 1
    for i, h in enumerate(['과세표준 초과', '기본 세액', '', '초과분 세율']):
        c = k.cell(r, 2 + i, h); c.font = f(10, True, INK2); c.border = Border(bottom=thin)
    r += 1
    t0 = r
    for lo, base, rate in TAX_TABLE:
        k.cell(r, 2, lo).number_format = '#,##0'
        k.cell(r, 3, base).number_format = '#,##0'
        k.cell(r, 5, rate).number_format = '0%'
        for col in (2, 3, 5):
            k.cell(r, col).font = f(10); k.cell(r, col).fill = PatternFill('solid', fgColor=BG)
        r += 1
    t1 = r - 1
    A['TLO'], A['TBASE'], A['TRATE'] = (f'기준값!$B${t0}:$B${t1}', f'기준값!$C${t0}:$C${t1}', f'기준값!$E${t0}:$E${t1}')
    k.cell(r, 2, '원문 표: 1,400만원 이하 6% / 84만원+1,400만원 초과분 15% / 624만원+5,000만원 초과분 24% / 1,536만원+8,800만원 초과분 35% / '
                 '3,706만원+1억5천만원 초과분 38% / 9,406만원+3억원 초과분 40% / 1억7,406만원+5억원 초과분 42% / 3억8,406만원+10억원 초과분 45%').font = f(9, color=INK2)
    r += 2
    # 재산 점수표
    k.cell(r, 2, '재산등급별 점수 — 국민건강보험법 시행령 별표 4 제3호(2026.10.1 시행판), 재산금액은 1억원을 뺀 뒤 금액').font = f(11, True); r += 1
    for i, h in enumerate(['등급', '초과(원)', '', '이하(원)', '점수']):
        c = k.cell(r, 2 + i, h); c.font = f(10, True, INK2); c.border = Border(bottom=thin)
    r += 1
    p0 = r
    lowers = [0] + [u * 10000 for u in PROP_UPPER_MAN]
    for i, pts in enumerate(PROP_POINTS):
        k.cell(r, 2, i + 1)
        k.cell(r, 3, lowers[i]).number_format = '#,##0'
        k.cell(r, 5, (PROP_UPPER_MAN[i] * 10000) if i < len(PROP_UPPER_MAN) else '그 이상').number_format = '#,##0'
        k.cell(r, 6, pts).number_format = '#,##0'
        for col in (2, 3, 5, 6):
            k.cell(r, col).font = f(10); k.cell(r, col).alignment = Alignment(horizontal='left')
        r += 1
    p1 = r - 1
    A['PLO'], A['PPTS'] = f'기준값!$C${p0}:$C${p1}', f'기준값!$F${p0}:$F${p1}'
    r += 1
    k.cell(r, 2, '확인 필요 — 원문으로 확인하지 못해 계산에 쓰지 않은 것').font = f(11, True, ORANGE); r += 1
    for i, h in enumerate(['항목', '', '', '확인한 것', '', '이 엑셀에서']):
        if h:
            c = k.cell(r, 2 + i, h); c.font = f(10, True, INK2)
    for col in range(2, 9):
        k.cell(r, col).border = Border(bottom=thin)
    r += 1
    for name, seen, here in UNCHECKED:
        k.cell(r, 2, name).font = f(10, True)
        k.cell(r, 5, seen).font = f(9, color=INK2)
        k.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
        k.cell(r, 7, here).font = f(9)
        k.cell(r, 8, '확인 필요').font = f(9, True, ORANGE)
        for col in (2, 5, 7, 8):
            k.cell(r, col).alignment = Alignment(wrap_text=True, vertical='top')
        k.row_dimensions[r].height = 62
        r += 1
    k.freeze_panes = 'C6'
    k.page_setup.orientation = 'landscape'
    k.page_setup.fitToWidth = 1; k.page_setup.fitToHeight = 0
    k.sheet_properties.pageSetUpPr.fitToPage = True

    # ── 판정 시트 ─────────────────────────────────────────────────
    s.sheet_view.showGridLines = False
    for col, w in zip('ABCDEF', (2, 40, 17, 12, 46, 2)):
        s.column_dimensions[col].width = w
    X = lambda key: A[key]  # noqa: E731
    I, DK, DE, ST, PEN, OTH, PROP = 'C12', 'C13', 'C14', 'C15', 'C16', 'C17', 'C18'

    def tax(x):
        idx = f'COUNTIF({X("TLO")},"<"&({x}))'
        return f'IF(({x})<=0,0,INDEX({X("TBASE")},{idx})+(({x})-INDEX({X("TLO")},{idx}))*INDEX({X("TRATE")},{idx}))'

    def premium(fin_cell, pen_cell, prop_cell):
        """지역가입자 건강보험료(장기요양 전, 원/월). fin_cell=건보에 넣는 금융소득(연)."""
        mon = f'(({fin_cell})+({pen_cell})*{X("HALF")})/12'
        inc = f'IF({mon}<={X("FLOORM")},{X("FLOOR")},{mon}*{X("HR")})'
        base = f'MAX(0,({prop_cell})-{X("PDED")})'
        prop = f'IF({base}<=0,0,INDEX({X("PPTS")},COUNTIF({X("PLO")},"<"&{base}))*{X("PT")})'
        return f'MIN(MAX({inc}+{prop},{X("FLOOR")}),{X("CAP")})'

    # 입력
    s['B11'] = '넣는 칸'; s['B11'].font = f(11, True)
    s['E11'] = '1년 금액, 세금 떼기 전 · 빈칸은 0'; s['E11'].font = f(9, color=INK2)
    inputs = [
        (12, '이자', ex['interest'], '원', '예금·적금·채권 이자'),
        (13, '배당 · 국내 상장주식', ex['div_kr'], '원', '배당가산 대상. 잘 모르면 아래 칸에 넣으세요(세금이 실제와 같거나 크게 나와요)'),
        (14, '배당 · ETF 분배금·해외주식 등', ex['div_etc'], '원', '국내 증권사 계좌에서 세금을 떼고 받은 것'),
        (15, '지금 건강보험', ex['status'], '', '직장가입자 가족의 피부양자인지, 지역가입자인지'),
        (16, '(선택) 공적연금·근로소득', ex['pension'], '원', '건보 계산에만 씀. 국민연금 등 1년 받은 총액. 사적연금은 넣지 않아요'),
        (17, '(선택) 그 밖의 종합소득금액', ex['other'], '원', '세금 계산에만 씀. 종합소득세 신고서의 소득금액(이자·배당 뺀 것)'),
        (18, '(선택) 재산세 과세표준 합계', ex['prop'], '원', '집·땅·건물. 재산세 고지서의 과세표준(시세보다 낮아요)'),
    ]
    for row, lab, val, unit, hint in inputs:
        s[f'B{row}'] = lab; s[f'B{row}'].font = f(11)
        c = s[f'C{row}']; c.value = val; c.font = f(11, True)
        c.number_format = '#,##0' if unit == '원' else '@'
        c.fill = PatternFill('solid', fgColor=BG)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        s[f'D{row}'] = unit; s[f'D{row}'].font = f(10, color=INK2)
        s[f'E{row}'] = hint; s[f'E{row}'].font = f(9, color=INK2)
        s[f'E{row}'].alignment = Alignment(wrap_text=True, vertical='center')
        s.row_dimensions[row].height = 30
    dv = DataValidation(type='list', formula1=f'"{DEP},{LOC}"', allow_blank=False)
    s.add_data_validation(dv); dv.add(ST)
    s[ST].alignment = Alignment(horizontal='left')

    # 근거 표
    rows = []  # (행, 항목, 식, 서식, 근거)
    R = {}

    def add(key, label, formula, fmt='#,##0', src='', bold=False):
        rows.append((key, label, formula, fmt, src, bold))

    N = lambda c: f'N({c})'  # noqa: E731  빈칸·글자를 0으로
    F_ = f'({N(I)}+{N(DK)}+{N(DE)})'
    add('h1', '① 금융소득 종합과세 — 소득세', None)
    add('F', '금융소득 합계', f'={F_}', src='이자 + 배당(받은 금액, 세전)')
    add('OVER', '2천만원을 넘는 금액', f'=MAX(0,{{F}}-{X("T")})', src='소득세법 제14조제3항제6호')
    add('GUB', '배당가산 대상 배당 가운데 넘는 부분', f'=MIN({N(DK)},{{OVER}})', src='이자 → 그 밖의 배당 → 국내 상장주식 배당 순으로 2천만원을 채운다(시행령 제116조의2)')
    add('G', '배당가산액', f'=ROUND({{GUB}}*{X("GU")},0)', src='소득세법 제17조제3항·제56조제4항')
    add('BASE1', '기본세율을 매기는 금액', f'=IF({{F}}>{X("T")},MAX(0,{{OVER}}+{{G}}+{N(OTH)}-{X("DED")}),"")',
        src='2천만원 초과 금융소득 + 배당가산액 + 다른 종합소득금액 − 소득공제')
    add('TAX1', '종합과세 시 세액 ①', f'=IF({{F}}>{X("T")},ROUND({X("T")}*{X("WI")}+{tax("{BASE1}")},0),"")',
        src='2천만원×14% + 위 금액×기본세율(제62조제1호, 국세상담센터 산식 ①)')
    other_base = f'MAX(0,{N(OTH)}-{X("DED")})'
    add('TAX2', '분리과세 시 세액 ②', f'=ROUND({N(I)}*{X("WI")}+({N(DK)}+{N(DE)})*{X("WD")}+{tax(other_base)},0)',
        src='금융소득×14% + (다른 종합소득금액 − 소득공제)×기본세율(제62조제2호, 산식 ②)')
    add('CR', '배당세액공제', f'=IF({{F}}>{X("T")},{{G}},0)', src='소득세법 제56조제1항')
    add('DET', '세금 기준 금액', f'=IF({{F}}>{X("T")},MAX({{TAX1}}-{{CR}},{{TAX2}}),{{TAX2}})', src='(① − 배당세액공제)와 ② 가운데 큰 것(제15조제2호·제62조)')
    add('ADD', '종합과세로 더 내는 소득세(연)', '={DET}-{TAX2}', src='세금 기준 금액 − ②. ②는 원천징수 14%와 다른 소득 세금이라 어차피 내는 돈', bold=True)
    add('h2', '② 피부양자 — 소득·재산 요건', None)
    add('FH', '건보 소득에 들어가는 금융소득', f'=IF({{F}}>{X("FT")},{{F}},0)', src='1천만원 이하면 0, 넘으면 전액(시행규칙 제44조제1항 단서)')
    add('SUMA', '합산소득 A — 금융소득을 모두 넣으면', f'={{F}}+{N(PEN)}', src='공단 누리집 "모든 소득(사업·금융·연금·근로기타소득 등)을 합하여"')
    add('SUMB', '합산소득 B — 1천만원 이하 금융소득을 빼면', f'={{FH}}+{N(PEN)}', src='소득월액 규칙(시행규칙 제44조제1항 단서)을 그대로 쓸 때')
    add('LIM', '소득 한도', f'=IF({N(PROP)}>{X("DEPP1")},{X("DEPINC2")},{X("DEPINC")})', src='재산 과표 5.4억원 이하 2천만원, 5.4억~9억원 1천만원(별표 1의2)')
    add('JA', '판정 A', f'=IF({N(PROP)}>{X("DEPP2")},"탈락",IF({{SUMA}}<={{LIM}},"유지","탈락"))', fmt='@', src='재산 과표 9억원 넘으면 탈락(별표 1의2 제2호)')
    add('JB', '판정 B', f'=IF({N(PROP)}>{X("DEPP2")},"탈락",IF({{SUMB}}<={{LIM}},"유지","탈락"))', fmt='@')
    add('JD', '피부양자 판정', '=IF({JA}={JB},{JA},"확인 필요")', fmt='@', src='A와 B가 다르면 공단(1577-1000)에 확인. 기준값 시트 "확인 필요" 첫 줄', bold=True)
    add('h3', '③ 지역가입자 — 월 건강보험료', None)
    add('INCY', '건보 소득(연)', f'={{FH}}+{N(PEN)}*{X("HALF")}', src='금융소득(1천만원 넘을 때) + 공적연금·근로소득의 50%')
    add('INCM', '소득월액', '={INCY}/12', src='연간 소득 ÷ 12(국민건강보험법 제71조제2항)')
    add('PI', '소득 보험료', f'=IF({{INCM}}<={X("FLOORM")},{X("FLOOR")},{{INCM}}*{X("HR")})', src='소득월액 × 7.19%. 28만원 이하면 하한 20,160원(시행규칙 제44조제3항, 고시 제3조)')
    add('PB', '재산 금액(1억원 공제 뒤)', f'=MAX(0,{N(PROP)}-{X("PDED")})', src='시행령 별표 4 제1호')
    add('PP', '재산 점수', f'=IF({{PB}}<=0,0,INDEX({X("PPTS")},COUNTIF({X("PLO")},"<"&{{PB}})))', src='시행령 별표 4 제3호 60등급')
    add('PR', '재산 보험료', f'={{PP}}*{X("PT")}', src='점수 × 211.5원(시행령 제44조제2항)')
    add('HI', '건강보험료', f'=MIN(MAX({{PI}}+{{PR}},{X("FLOOR")}),{X("CAP")})', src='소득 + 재산(국민건강보험법 제69조제5항), 하한 20,160원·상한 4,591,740원(제69조제6항, 고시)')
    add('LT', '장기요양보험료', f'={{HI}}*{X("LTC")}', src='건강보험료 × 13.14%(0.9448÷7.19를 넷째자리까지, 노인장기요양보험법 제9조)')
    add('TOT', '지역가입자 월 보험료 합계', '=ROUND({HI}+{LT},0)', src='건강보험료 + 장기요양보험료', bold=True)
    add('TOT0', '금융소득이 없을 때 월 보험료', f'=ROUND({premium("0", N(PEN), N(PROP))}*(1+{X("LTC")}),0)', src='같은 식에 금융소득 0')
    add('DIFF', '금융소득 때문에 늘어난 월 보험료', '={TOT}-{TOT0}', bold=True)

    s['B20'] = '근거 표'; s['B20'].font = f(11, True)
    s['E20'] = '근거'; s['E20'].font = f(10, True, INK2)
    r = 21
    for key, *_ in rows:
        R[key] = f'C{r}'
        r += 1
    r = 21
    for key, label, formula, fmt, src, bold in rows:
        if formula is None:
            s[f'B{r}'] = label; s[f'B{r}'].font = f(11, True, ORANGE)
            for col in 'BCDE':
                s[f'{col}{r}'].border = Border(bottom=thin)
        else:
            s[f'B{r}'] = label; s[f'B{r}'].font = f(10, bold)
            c = s[f'C{r}']; c.value = formula.format(**R); c.font = f(10, bold)
            c.number_format = fmt
            c.alignment = Alignment(horizontal='right')
            if fmt == '#,##0':
                s[f'D{r}'] = '원'; s[f'D{r}'].font = f(9, color=INK2)
            s[f'E{r}'] = src; s[f'E{r}'].font = f(9, color=INK2)
            s[f'E{r}'].alignment = Alignment(wrap_text=True, vertical='center')
        r += 1
    last = r - 1
    s['D' + R['PP'][1:]] = '점'

    # 큰 숫자·타일 (위)
    ADD, JD, TOT, TOT0, DIFF, JA, JB, F = (R[x] for x in ('ADD', 'JD', 'TOT', 'TOT0', 'DIFF', 'JA', 'JB', 'F'))
    won = lambda c: f'TEXT({c},"#,##0")&"원"'  # noqa: E731
    dep = f'{ST}="{DEP}"'
    s['B2'] = (f'=TRIM(IF(AND({dep},{JD}="확인 필요"),"피부양자 판정이 갈려요: 1천만원 이하 금융소득을 소득에 넣는지 공단 안내에 없어요(기준값 시트). ","")'
               f'&IF(AND({N(DE)}>0,{ADD}>0),"해외주식 배당이 있다면 외국에서 낸 세금 공제를 넣지 않아 더 내는 소득세가 실제보다 크게 나올 수 있어요. ",""))')
    s['B2'].font = f(10, True, ORANGE)
    s['B3'] = f'=IF({dep},IF({JD}="유지","피부양자 유지 · 건보료",IF({JD}="탈락","피부양자에서 빠지면 지역가입자로 내는 건보료","피부양자 판정 확인 필요")),"금융소득 때문에 붙는 건보료")'
    s['B3'].font = f(11, color=INK2)
    s.merge_cells('B4:E6')
    s['B4'] = (f'=IF({dep},IF({JD}="유지","월 0원",IF({JD}="탈락","월 "&{won(TOT)},"유지 또는 월 "&{won(TOT)})),'
               f'"월 "&{won(DIFF)})')
    s['B4'].font = f(36, True, ORANGE)
    s['B4'].alignment = Alignment(vertical='center')
    s['B7'] = (f'=IF({dep},IF({JD}="유지","지금 금액이면 건보료가 붙지 않아요",IF({JD}="탈락","탈락하면 지역가입자 월 "&{won(TOT)}&" · 금융소득이 없어도 월 "&{won(TOT0)},'
               f'"공단(1577-1000) 확인 필요 · 탈락하면 지역가입자 월 "&{won(TOT)})),'
               f'IF({DIFF}<=0,"금융소득 때문에 늘어나는 건보료는 없어요 · 전체 월 "&{won(TOT)},'
               f'"지금 이 금액이면 건보료가 월 "&{won(DIFF)}&" 붙어요 · 전체 월 "&{won(TOT)}))')
    s['B7'].font = f(11, color=INK2)
    s['B8'] = '=IF({0}<=0,"종합과세로 더 내는 소득세는 없어요","종합과세로 소득세를 1년에 "&{1}&" 더 내요")'.format(ADD, won(ADD))
    s['B8'].font = f(11, color=INK2)
    for col, lab, ref in [('B', '종합과세로 더 내는 소득세(연)', f'={ADD}'), ('C', '피부양자 판정', f'={JD}'), ('E', '지역가입자면 월 건보료', f'={TOT}')]:
        s[f'{col}9'] = lab; s[f'{col}9'].font = f(10, color=INK2)
        c = s[f'{col}10']; c.value = ref; c.font = f(14, True); c.alignment = Alignment(horizontal='left')
        c.number_format = '#,##0"원"'
        c.border = Border(bottom=thin)
    s.merge_cells('C9:D9'); s.merge_cells('C10:D10')
    s[f'B{last + 2}'] = '2026년 요율 기준 계산 도구예요. 투자 권유가 아니고, 정확한 금액은 국세청·국민건강보험공단에서 확인하세요.'
    s[f'B{last + 2}'].font = f(9, color=INK2)
    s[f'B{last + 3}'] = '파이어맵 건보료 계산기와 같은 줄기 · firemap.kr/health-insurance'
    s[f'B{last + 3}'].hyperlink = 'https://firemap.kr/health-insurance?utm_source=xlsx&utm_medium=sheet&utm_campaign=fin1'
    s[f'B{last + 3}'].font = f(9, color=INK2, underline='single')
    s.print_area = f'A1:F{last + 3}'
    s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True

    # ── 설명 시트 ─────────────────────────────────────────────────
    g.sheet_view.showGridLines = False
    g.column_dimensions['A'].width = 2
    g.column_dimensions['B'].width = 104
    lines = [
        ('쓰는 법', True),
        ('1. 「판정」 시트 회색 칸에 1년 이자·배당(세금 떼기 전)을 넣고, 지금 건강보험 자격을 고릅니다.', False),
        ('2. 국민연금 같은 공적연금, 다른 종합소득, 재산세 과세표준이 있으면 (선택) 칸에 넣어요. 없으면 비워 두세요.', False),
        ('3. 맨 위 큰 숫자가 건보료, 그 아래 줄이 종합과세로 더 내는 소득세입니다. 계산 과정은 「근거 표」에 한 줄씩 있어요.', False),
        ('', False),
        ('무엇을 계산하나', True),
        ('· 소득세: 금융소득이 2천만원을 넘으면 (① 종합과세 시 세액 − 배당세액공제)와 (② 분리과세 시 세액) 가운데 큰 것을 냅니다. ②와의 차이를 「더 내는 소득세」로 보여 줍니다.', False),
        ('· 피부양자: 소득 합계 2천만원(재산세 과세표준 5.4억~9억원이면 1천만원), 재산세 과세표준 9억원 기준으로 봅니다.', False),
        ('· 지역가입자: (소득월액 × 7.19% 또는 하한 20,160원) + (재산 점수 × 211.5원), 여기에 장기요양보험료를 더합니다.', False),
        ('', False),
        ('언제 반영되나', True),
        ('· 이자·배당 같은 소득은 다음 해 11월부터 그다음 해 10월까지 건보료에 들어갑니다(국민건강보험법 시행령 제41조제3항). 예: 2026년 이자·배당 → 2027년 11월~2028년 10월.', False),
        ('· 종합소득세는 다음 해 5월에 신고합니다.', False),
        ('', False),
        ('넣지 않은 것', True),
        ('· 지방소득세, 외국에서 낸 세금 공제, 표준세액공제 같은 다른 세액공제, 고배당기업 배당 분리과세 특례.', False),
        ('· 국내에서 세금을 떼지 않은 이자·배당(해외 계좌로 직접 받은 것 등). 이런 소득은 2천만원 이하여도 종합과세 대상입니다(국세상담센터).', False),
        ('· 사업소득·주택임대소득이 있으면 피부양자 판정이 달라집니다(별표 1의2 제1호나목). 이 파일은 이 경우를 판정하지 않아요.', False),
        ('· 건보료는 세대 단위로 매깁니다. 이 파일은 한 사람 몫만 계산하고, 전월세·주택금융부채 공제·자동차는 넣지 않았습니다.', False),
        ('· 확인하지 못한 기준은 「기준값」 시트 아래 「확인 필요」에 모았습니다. 계산에는 쓰지 않았어요.', False),
        ('', False),
        ('알림', True),
        ('· 이 파일은 계산 도구이며 투자 조언이 아닙니다. 특정 금융상품을 추천하지 않습니다.', False),
        ('· 기준일 2026-10-12. 보험료율·세율이 바뀌면 「기준값」 시트의 값만 고치면 됩니다.', False),
        ('firemap.kr', False),
    ]
    for i, (t, bold) in enumerate(lines):
        c = g.cell(2 + i, 2, t); c.font = f(12 if bold else 11, bold)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    link = g.cell(2 + len(lines) - 1, 2)
    link.hyperlink = 'https://firemap.kr/?utm_source=xlsx&utm_medium=guide&utm_campaign=fin1'
    link.font = f(11, color=INK2, underline='single')

    wb.active = 0
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path, R, A


if __name__ == '__main__':
    p, R, A = build()
    print(p)
