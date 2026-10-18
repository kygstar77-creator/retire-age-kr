# 회원 자료 2호 '내 연봉 상위 몇 %' 엑셀 (firemap-write, 2026-10-18)
#   py -3.12 make_xlsx.py   → out/연봉_상위몇퍼센트_국세청2024.xlsx
# 원자료: 국세청 「근로소득 백분위(천분위) 자료」 공공데이터포털 15082063(2024년 귀속, 수정 2026-05-08)
#   raw_nts_pct_utf8.csv = longform/ep/N-1/raw/nts_pct_utf8.csv 복사본(인원 명, 금액 억원, 줄 = 구간 합계)
# 구간 평균만 있고 경계값은 없다 → 파일은 "어느 두 구간 평균 사이인지"만 알려 준다.
import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
FILE = os.path.join(HERE, 'out', '연봉_상위몇퍼센트_국세청2024.xlsx')
SRC_URL = 'https://www.data.go.kr/data/15082063/fileData.do'


def load():
    rows = []
    for r in csv.reader(open(os.path.join(HERE, 'raw_nts_pct_utf8.csv'), encoding='utf-8')):
        if not r or '상위' not in r[0]:
            continue
        rows.append(dict(k=r[0].strip(), n=int(r[1]), g=int(r[2])))
    return rows  # 위(상위 0.1%) → 아래(상위 100%)


ROWS = load()
FIRST = 4                        # 원자료 표 첫 줄
LAST = FIRST + len(ROWS) - 1     # 마지막 줄


def build():
    wb = Workbook()
    bold = Font(bold=True)
    big = Font(bold=True, size=16)
    inp = PatternFill('solid', fgColor='FFF2CC')
    out = PatternFill('solid', fgColor='E2EFDA')
    thin = Side(style='thin', color='BBBBBB')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ── 원자료(아래 → 위 순서로 둔다: MATCH 근사 검색은 오름차순이어야 한다)
    d = wb.active
    d.title = '원자료'
    d['A1'] = '국세청 근로소득 백분위(천분위) 자료 · 2024년 귀속(2025년 연말정산) · 공공데이터포털 15082063(수정 2026-05-08)'
    d['A2'] = SRC_URL
    hdr = ['구간', '인원(명)', '총급여 합계(억원)', '구간 평균 총급여(만원)', '이 구간보다 위에 있는 사람 비율']
    for i, h in enumerate(hdr):
        c = d.cell(row=3, column=i + 1, value=h); c.font = bold; c.border = box
    for i, r in enumerate(reversed(ROWS)):
        y = FIRST + i
        d.cell(row=y, column=1, value=r['k'])
        d.cell(row=y, column=2, value=r['n']).number_format = '#,##0'
        d.cell(row=y, column=3, value=r['g']).number_format = '#,##0'
        d.cell(row=y, column=4, value=f'=ROUND(C{y}*10000/B{y},0)').number_format = '#,##0'
        d.cell(row=y, column=5, value=f'=SUM(B{y + 1}:B{LAST + 1})/SUM($B${FIRST}:$B${LAST})').number_format = '0.0%'
    d.cell(row=LAST + 2, column=1, value='합계').font = bold
    d.cell(row=LAST + 2, column=2, value=f'=SUM(B{FIRST}:B{LAST})').number_format = '#,##0'
    d.cell(row=LAST + 2, column=3, value=f'=SUM(C{FIRST}:C{LAST})').number_format = '#,##0'
    d.cell(row=LAST + 2, column=4, value=f'=ROUND(C{LAST + 2}*10000/B{LAST + 2},0)').number_format = '#,##0'
    for col, w in zip('ABCDE', (14, 14, 18, 22, 28)):
        d.column_dimensions[col].width = w
    d.freeze_panes = 'A4'

    # ── 내 연봉 넣기
    s = wb.create_sheet('내 연봉 넣기', 0)
    s['A1'] = '내 연봉은 상위 몇 % 근처일까'; s['A1'].font = big
    s['A2'] = '국세청 2024년 귀속 근로소득 자료(연말정산한 근로자 2,108만명)로 봅니다.'
    s['A4'] = '① 1년 총급여(만원)'; s['A4'].font = bold
    s['C4'] = 5000; s['C4'].fill = inp; s['C4'].number_format = '#,##0'; s['C4'].border = box
    s['D4'] = '← 노란 칸에 숫자만 넣으세요. 원천징수영수증의 「총급여」(비과세 뺀 세전 금액)'
    dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='10000000', allow_blank=False)
    dv.error = '0 이상 숫자(만원)만 넣을 수 있어요'; s.add_data_validation(dv); dv.add('C4')

    rng_k = f"원자료!$A${FIRST}:$A${LAST}"
    rng_a = f"원자료!$D${FIRST}:$D${LAST}"
    s['A6'] = '② 바로 아래 구간'; s['A6'].font = bold
    s['C6'] = f'=IF(C4<MIN({rng_a}),"맨 아래 구간 평균보다 적음",INDEX({rng_k},MATCH(C4,{rng_a},1)))'
    s['A7'] = '   그 구간 평균(만원)'
    s['C7'] = f'=IF(C4<MIN({rng_a}),"",INDEX({rng_a},MATCH(C4,{rng_a},1)))'
    s['A8'] = '③ 바로 위 구간'; s['A8'].font = bold
    s['C8'] = (f'=IF(C4<MIN({rng_a}),INDEX({rng_k},1),IF(C4>=MAX({rng_a}),"위 구간 없음(상위 0.1% 평균 이상)",'
               f'INDEX({rng_k},MATCH(C4,{rng_a},1)+1)))')
    s['A9'] = '   그 구간 평균(만원)'
    s['C9'] = (f'=IF(C4<MIN({rng_a}),INDEX({rng_a},1),IF(C4>=MAX({rng_a}),"",'
               f'INDEX({rng_a},MATCH(C4,{rng_a},1)+1)))')
    for a in ('C6', 'C7', 'C8', 'C9'):
        s[a].fill = out; s[a].border = box; s[a].number_format = '#,##0'
    s['A11'] = '한 줄로'; s['A11'].font = bold
    s['A12'] = ('=IF(C4>=MAX(' + rng_a + '),"총급여 "&TEXT(C4,"#,##0")&"만원은 상위 0.1% 구간 평균("&TEXT(MAX(' + rng_a + '),"#,##0")&"만원)과 같거나 많아요.",'
                'IF(C4<MIN(' + rng_a + '),"총급여 "&TEXT(C4,"#,##0")&"만원은 맨 아래 구간 평균("&TEXT(MIN(' + rng_a + '),"#,##0")&"만원)보다 적어요.",'
                '"총급여 "&TEXT(C4,"#,##0")&"만원은 「"&C6&"」 구간 평균("&TEXT(C7,"#,##0")&"만원)과 「"&C8&"」 구간 평균("&TEXT(C9,"#,##0")&"만원) 사이예요."))')
    s['A12'].font = Font(bold=True, size=12)

    s['A14'] = '읽는 법'; s['A14'].font = bold
    notes = [
        '· 국세청 자료는 구간마다 「합계」만 있어요. 그래서 정확한 경계값(상위 10%가 몇 만원부터인지)은 못 짚고, 두 구간 평균 사이로만 말합니다.',
        '· 예) 「상위 10%」 구간 = 위에서 9%와 10% 사이에 있는 약 21만명입니다. 「상위 10%」 구간 평균보다 많으면 위에서 10% 안쪽일 가능성이 큽니다.',
        '· 연말정산 자료라 중간에 입사·퇴사한 사람도 들어 있어요. 맨 아래 1% 구간 평균은 21만원입니다.',
        '· 총급여는 비과세(식대 등)를 뺀 세전 금액입니다. 사업·임대·금융 소득은 들어 있지 않아요.',
        '· 2024년 한 해(2025년 연말정산) 자료입니다. 다음 해 자료가 나오면 원자료 시트만 바꾸면 됩니다.',
        '· 투자·이직 권유가 아닙니다. 공개 통계로 내 위치를 보는 표입니다.',
    ]
    for i, t in enumerate(notes):
        s.cell(row=15 + i, column=1, value=t)
    s['A22'] = '출처: 국세청 근로소득 백분위(천분위) 자료, 공공데이터포털 ' + SRC_URL
    s['A23'] = '만든 곳: 파이어맵 카페(cafe.naver.com/firemap) · 검산 기록은 카페 글에 있습니다.'
    s.column_dimensions['A'].width = 22; s.column_dimensions['B'].width = 2
    s.column_dimensions['C'].width = 30; s.column_dimensions['D'].width = 60

    # ── 주요 구간 한눈에(값은 원자료 시트 수식을 그대로 가리킨다)
    t = wb.create_sheet('주요 구간', 1)
    t['A1'] = '주요 구간 평균 총급여(만원)'; t['A1'].font = big
    t['A3'] = '구간'; t['B3'] = '구간 평균(만원)'; t['C3'] = '뜻'
    for c in ('A3', 'B3', 'C3'):
        t[c].font = bold; t[c].border = box
    keys = ['상위 0.1%', '상위 1.0%', '상위 5%', '상위 10%', '상위 20%', '상위 30%', '상위 50%', '상위 51%', '상위 70%', '상위 90%']
    order = [r['k'] for r in reversed(ROWS)]
    for i, k in enumerate(keys):
        y = 4 + i
        ref = FIRST + order.index(k)
        t.cell(row=y, column=1, value=k)
        t.cell(row=y, column=2, value=f'=원자료!D{ref}').number_format = '#,##0'
        lo = k.replace('상위 ', '')
        t.cell(row=y, column=3, value=f'위에서 {lo} 언저리 사람들의 평균')
    t['A15'] = '「상위 50%」·「상위 51%」 두 구간 평균 사이가 한가운데(중간) 근처입니다.'
    t.column_dimensions['A'].width = 14; t.column_dimensions['B'].width = 16; t.column_dimensions['C'].width = 36

    os.makedirs(os.path.dirname(FILE), exist_ok=True)
    wb.save(FILE)
    return FILE


if __name__ == '__main__':
    print(build())
